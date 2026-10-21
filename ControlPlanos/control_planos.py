"""Control de Planos - registro de entrega y devolución de planos.

Guarda los datos en una base SQLite (control_planos.db) ubicada junto al
programa y permite exportar la planilla a Excel.
"""

import os
import sqlite3
import sys
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

APP_TITULO = "Control de Planos"
FORMATO_FECHA = "%d/%m/%Y %H:%M"


def carpeta_programa():
    """Carpeta donde está el .exe (o el .py) para guardar la base de datos."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def ruta_base_datos():
    carpeta = carpeta_programa()
    if os.access(carpeta, os.W_OK):
        return os.path.join(carpeta, "control_planos.db")
    # Si la carpeta no permite escribir (ej. Archivos de programa) usamos AppData.
    alternativa = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), "ControlPlanos")
    os.makedirs(alternativa, exist_ok=True)
    return os.path.join(alternativa, "control_planos.db")


def formatear_fecha(valor):
    if not valor:
        return ""
    return datetime.fromisoformat(valor).strftime(FORMATO_FECHA)


# --------------------------------------------------------------------------
# Base de datos
# --------------------------------------------------------------------------
class BaseDatos:
    def __init__(self, ruta):
        self.ruta = ruta
        self.con = sqlite3.connect(ruta)
        self.con.row_factory = sqlite3.Row
        self.con.execute("PRAGMA foreign_keys = ON")
        self._crear_tablas()

    def _crear_tablas(self):
        self.con.executescript(
            """
            CREATE TABLE IF NOT EXISTS operadores (
                id      INTEGER PRIMARY KEY AUTOINCREMENT,
                nombre  TEXT NOT NULL UNIQUE COLLATE NOCASE
            );
            CREATE TABLE IF NOT EXISTS planos (
                id                INTEGER PRIMARY KEY AUTOINCREMENT,
                codigo            TEXT NOT NULL,
                contiene_muestra  INTEGER NOT NULL,
                nombre            TEXT NOT NULL,
                sector            TEXT NOT NULL,
                operador          TEXT NOT NULL,
                fecha_entrega     TEXT NOT NULL,
                fecha_devolucion  TEXT
            );
            """
        )
        self.con.commit()

    # ---- operadores ----
    def operadores(self):
        return [f["nombre"] for f in self.con.execute("SELECT nombre FROM operadores ORDER BY nombre")]

    def agregar_operador(self, nombre):
        self.con.execute("INSERT INTO operadores (nombre) VALUES (?)", (nombre,))
        self.con.commit()

    def eliminar_operador(self, nombre):
        self.con.execute("DELETE FROM operadores WHERE nombre = ?", (nombre,))
        self.con.commit()

    # ---- planos ----
    def sectores(self):
        return [f["sector"] for f in self.con.execute("SELECT DISTINCT sector FROM planos ORDER BY sector")]

    def registrar_entrega(self, codigo, contiene_muestra, nombre, sector, operador):
        cur = self.con.execute(
            """INSERT INTO planos (codigo, contiene_muestra, nombre, sector, operador, fecha_entrega)
               VALUES (?, ?, ?, ?, ?, ?)""",
            (codigo, 1 if contiene_muestra else 0, nombre, sector, operador,
             datetime.now().isoformat(timespec="seconds")),
        )
        self.con.commit()
        return cur.lastrowid

    def registrar_devolucion(self, id_plano):
        self.con.execute(
            "UPDATE planos SET fecha_devolucion = ? WHERE id = ? AND fecha_devolucion IS NULL",
            (datetime.now().isoformat(timespec="seconds"), id_plano),
        )
        self.con.commit()

    def anular_devolucion(self, id_plano):
        self.con.execute("UPDATE planos SET fecha_devolucion = NULL WHERE id = ?", (id_plano,))
        self.con.commit()

    def eliminar_plano(self, id_plano):
        self.con.execute("DELETE FROM planos WHERE id = ?", (id_plano,))
        self.con.commit()

    def buscar(self, texto="", estado="Todos"):
        sql = "SELECT * FROM planos WHERE 1 = 1"
        params = []
        if texto:
            sql += " AND (codigo LIKE ? OR nombre LIKE ? OR sector LIKE ? OR operador LIKE ?)"
            params += [f"%{texto}%"] * 4
        if estado == "Pendientes":
            sql += " AND fecha_devolucion IS NULL"
        elif estado == "Devueltos":
            sql += " AND fecha_devolucion IS NOT NULL"
        sql += " ORDER BY fecha_entrega DESC, id DESC"
        return self.con.execute(sql, params).fetchall()


# --------------------------------------------------------------------------
# Exportación a Excel
# --------------------------------------------------------------------------
ENCABEZADOS = ["Código de pieza", "Contiene muestra", "Nombre de la pieza", "Sector",
               "Operador", "Fecha de entrega", "Fecha de devolución", "Estado"]


def fila_para_mostrar(f):
    return (
        f["codigo"],
        "Sí" if f["contiene_muestra"] else "No",
        f["nombre"],
        f["sector"],
        f["operador"],
        formatear_fecha(f["fecha_entrega"]),
        formatear_fecha(f["fecha_devolucion"]),
        "Devuelto" if f["fecha_devolucion"] else "Entregado",
    )


def exportar_excel(filas, ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Planos"
    ws.append(ENCABEZADOS)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="305496")
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for f in filas:
        valores = list(fila_para_mostrar(f))
        # Fechas como fechas reales de Excel para poder ordenar/filtrar.
        valores[5] = datetime.fromisoformat(f["fecha_entrega"])
        valores[6] = datetime.fromisoformat(f["fecha_devolucion"]) if f["fecha_devolucion"] else None
        ws.append(valores)
        for col in (6, 7):
            ws.cell(row=ws.max_row, column=col).number_format = "dd/mm/yyyy hh:mm"

    anchos = [18, 16, 34, 20, 22, 18, 20, 12]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(ruta)


# --------------------------------------------------------------------------
# Interfaz
# --------------------------------------------------------------------------
class VentanaOperadores(tk.Toplevel):
    def __init__(self, padre, db, al_cerrar):
        super().__init__(padre)
        self.db = db
        self.al_cerrar = al_cerrar
        self.title("Operadores")
        self.geometry("360x380")
        self.transient(padre)
        self.grab_set()

        marco = ttk.Frame(self, padding=10)
        marco.pack(fill="both", expand=True)

        ttk.Label(marco, text="Nuevo operador:").pack(anchor="w")
        fila = ttk.Frame(marco)
        fila.pack(fill="x", pady=(2, 8))
        self.nombre = ttk.Entry(fila)
        self.nombre.pack(side="left", fill="x", expand=True)
        self.nombre.bind("<Return>", lambda e: self.agregar())
        ttk.Button(fila, text="Agregar", command=self.agregar).pack(side="left", padx=(6, 0))

        self.lista = tk.Listbox(marco, height=12)
        self.lista.pack(fill="both", expand=True)
        ttk.Button(marco, text="Eliminar seleccionado", command=self.eliminar).pack(fill="x", pady=(8, 0))

        self.protocol("WM_DELETE_WINDOW", self.cerrar)
        self.refrescar()
        self.nombre.focus_set()

    def refrescar(self):
        self.lista.delete(0, "end")
        for nombre in self.db.operadores():
            self.lista.insert("end", nombre)

    def agregar(self):
        nombre = self.nombre.get().strip()
        if not nombre:
            return
        try:
            self.db.agregar_operador(nombre)
        except sqlite3.IntegrityError:
            messagebox.showwarning(APP_TITULO, f"El operador '{nombre}' ya existe.", parent=self)
            return
        self.nombre.delete(0, "end")
        self.refrescar()

    def eliminar(self):
        sel = self.lista.curselection()
        if not sel:
            return
        nombre = self.lista.get(sel[0])
        if messagebox.askyesno(APP_TITULO, f"¿Eliminar al operador '{nombre}' de la lista?\n"
                               "(Los registros ya cargados no se modifican)", parent=self):
            self.db.eliminar_operador(nombre)
            self.refrescar()

    def cerrar(self):
        self.al_cerrar()
        self.destroy()


class App(tk.Tk):
    def __init__(self, db):
        super().__init__()
        self.db = db
        self.title(APP_TITULO)
        self.geometry("1150x650")
        self.minsize(900, 500)

        estilo = ttk.Style(self)
        if "vista" in estilo.theme_names():
            estilo.theme_use("vista")
        estilo.configure("Treeview", rowheight=24)
        estilo.configure("Accion.TButton", font=("Segoe UI", 10, "bold"))

        self._armar_formulario()
        self._armar_listado()
        self.refrescar_combos()
        self.refrescar_listado()
        self.codigo.focus_set()

    # ---- formulario de alta ----
    def _armar_formulario(self):
        marco = ttk.LabelFrame(self, text=" Nueva entrega de plano ", padding=10)
        marco.pack(fill="x", padx=10, pady=(10, 5))

        ttk.Label(marco, text="Código de pieza:").grid(row=0, column=0, sticky="w")
        self.codigo = ttk.Entry(marco, width=20)
        self.codigo.grid(row=1, column=0, sticky="we", padx=(0, 10))

        ttk.Label(marco, text="Nombre de la pieza:").grid(row=0, column=1, sticky="w")
        self.nombre = ttk.Entry(marco, width=35)
        self.nombre.grid(row=1, column=1, sticky="we", padx=(0, 10))

        ttk.Label(marco, text="¿Contiene muestra?").grid(row=0, column=2, sticky="w")
        self.muestra = tk.StringVar(value="No")
        fila_muestra = ttk.Frame(marco)
        fila_muestra.grid(row=1, column=2, sticky="w", padx=(0, 10))
        ttk.Radiobutton(fila_muestra, text="Sí", value="Sí", variable=self.muestra).pack(side="left")
        ttk.Radiobutton(fila_muestra, text="No", value="No", variable=self.muestra).pack(side="left", padx=(6, 0))

        ttk.Label(marco, text="Sector al que se dirige:").grid(row=0, column=3, sticky="w")
        self.sector = ttk.Combobox(marco, width=22)
        self.sector.grid(row=1, column=3, sticky="we", padx=(0, 10))

        ttk.Label(marco, text="Operador:").grid(row=0, column=4, sticky="w")
        fila_op = ttk.Frame(marco)
        fila_op.grid(row=1, column=4, sticky="we")
        self.operador = ttk.Combobox(fila_op, width=22, state="readonly")
        self.operador.pack(side="left", fill="x", expand=True)
        ttk.Button(fila_op, text="…", width=3, command=self.abrir_operadores).pack(side="left", padx=(4, 0))

        ttk.Button(marco, text="Registrar entrega", style="Accion.TButton",
                   command=self.registrar_entrega).grid(row=1, column=5, sticky="e", padx=(15, 0))

        for col, peso in ((0, 1), (1, 2), (3, 1), (4, 1)):
            marco.columnconfigure(col, weight=peso)

        for w in (self.codigo, self.nombre, self.sector):
            w.bind("<Return>", lambda e: self.registrar_entrega())

    # ---- listado ----
    def _armar_listado(self):
        barra = ttk.Frame(self, padding=(10, 5))
        barra.pack(fill="x")

        ttk.Label(barra, text="Buscar:").pack(side="left")
        self.busqueda = tk.StringVar()
        self.busqueda.trace_add("write", lambda *a: self.refrescar_listado())
        ttk.Entry(barra, textvariable=self.busqueda, width=30).pack(side="left", padx=(4, 12))

        ttk.Label(barra, text="Mostrar:").pack(side="left")
        self.estado = ttk.Combobox(barra, values=["Todos", "Pendientes", "Devueltos"], width=12, state="readonly")
        self.estado.set("Todos")
        self.estado.bind("<<ComboboxSelected>>", lambda e: self.refrescar_listado())
        self.estado.pack(side="left", padx=(4, 12))

        ttk.Button(barra, text="Exportar a Excel", command=self.exportar).pack(side="right")
        ttk.Button(barra, text="Eliminar registro", command=self.eliminar).pack(side="right", padx=4)
        ttk.Button(barra, text="Anular devolución", command=self.anular_devolucion).pack(side="right")
        ttk.Button(barra, text="Marcar como DEVUELTO", style="Accion.TButton",
                   command=self.marcar_devuelto).pack(side="right", padx=4)

        marco = ttk.Frame(self, padding=(10, 0, 10, 0))
        marco.pack(fill="both", expand=True)
        columnas = ("codigo", "muestra", "nombre", "sector", "operador", "entrega", "devolucion", "estado")
        anchos = (120, 120, 220, 130, 140, 130, 140, 90)
        self.tabla = ttk.Treeview(marco, columns=columnas, show="headings", selectmode="extended")
        for col, titulo, ancho in zip(columnas, ENCABEZADOS, anchos):
            self.tabla.heading(col, text=titulo)
            self.tabla.column(col, width=ancho, anchor="center" if col in ("muestra", "estado", "entrega", "devolucion") else "w")
        self.tabla.tag_configure("pendiente", background="#FFF4CE")
        self.tabla.tag_configure("devuelto", background="#E2EFDA")
        scroll = ttk.Scrollbar(marco, orient="vertical", command=self.tabla.yview)
        self.tabla.configure(yscrollcommand=scroll.set)
        self.tabla.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        self.tabla.bind("<Double-1>", lambda e: self.marcar_devuelto())

        self.estado_barra = tk.StringVar()
        ttk.Label(self, textvariable=self.estado_barra, anchor="w", padding=(10, 4)).pack(fill="x")

    # ---- acciones ----
    def refrescar_combos(self):
        self.operador["values"] = self.db.operadores()
        if self.operador.get() not in self.operador["values"]:
            self.operador.set("")
        self.sector["values"] = self.db.sectores()

    def refrescar_listado(self):
        self.tabla.delete(*self.tabla.get_children())
        filas = self.db.buscar(self.busqueda.get().strip(), self.estado.get())
        pendientes = 0
        for f in filas:
            tag = "devuelto" if f["fecha_devolucion"] else "pendiente"
            pendientes += tag == "pendiente"
            self.tabla.insert("", "end", iid=str(f["id"]), values=fila_para_mostrar(f), tags=(tag,))
        self.estado_barra.set(f"{len(filas)} registro(s) mostrados  ·  {pendientes} sin devolver  ·  "
                              f"Base de datos: {self.db.ruta}")

    def abrir_operadores(self):
        VentanaOperadores(self, self.db, self.refrescar_combos)

    def registrar_entrega(self):
        codigo = self.codigo.get().strip()
        nombre = self.nombre.get().strip()
        sector = self.sector.get().strip()
        operador = self.operador.get().strip()
        faltan = [n for n, v in (("Código de pieza", codigo), ("Nombre de la pieza", nombre),
                                 ("Sector", sector), ("Operador", operador)) if not v]
        if faltan:
            extra = "\n\nUse el botón «…» para cargar operadores." if not self.db.operadores() else ""
            messagebox.showwarning(APP_TITULO, "Complete los campos:\n- " + "\n- ".join(faltan) + extra)
            return
        self.db.registrar_entrega(codigo, self.muestra.get() == "Sí", nombre, sector, operador)
        self.codigo.delete(0, "end")
        self.nombre.delete(0, "end")
        self.muestra.set("No")
        self.refrescar_combos()
        self.refrescar_listado()
        self.codigo.focus_set()

    def _seleccion(self):
        sel = self.tabla.selection()
        if not sel:
            messagebox.showinfo(APP_TITULO, "Seleccione uno o más registros de la lista.")
        return [int(i) for i in sel]

    def marcar_devuelto(self):
        ids = self._seleccion()
        if not ids:
            return
        for i in ids:
            self.db.registrar_devolucion(i)
        self.refrescar_listado()

    def anular_devolucion(self):
        ids = self._seleccion()
        if ids and messagebox.askyesno(APP_TITULO, "¿Quitar la fecha de devolución de los registros seleccionados?"):
            for i in ids:
                self.db.anular_devolucion(i)
            self.refrescar_listado()

    def eliminar(self):
        ids = self._seleccion()
        if ids and messagebox.askyesno(APP_TITULO, f"¿Eliminar definitivamente {len(ids)} registro(s)?"):
            for i in ids:
                self.db.eliminar_plano(i)
            self.refrescar_combos()
            self.refrescar_listado()

    def exportar(self):
        filas = self.db.buscar(self.busqueda.get().strip(), self.estado.get())
        if not filas:
            messagebox.showinfo(APP_TITULO, "No hay registros para exportar.")
            return
        ruta = filedialog.asksaveasfilename(
            title="Exportar a Excel", defaultextension=".xlsx",
            initialfile=f"Planos_{datetime.now():%Y-%m-%d}.xlsx",
            filetypes=[("Libro de Excel", "*.xlsx")])
        if not ruta:
            return
        try:
            exportar_excel(filas, ruta)
        except PermissionError:
            messagebox.showerror(APP_TITULO, "No se pudo guardar. ¿El archivo está abierto en Excel?")
            return
        if messagebox.askyesno(APP_TITULO, f"Exportado {len(filas)} registro(s) a:\n{ruta}\n\n¿Abrir el archivo?"):
            if hasattr(os, "startfile"):
                os.startfile(ruta)


def main():
    db = BaseDatos(ruta_base_datos())
    App(db).mainloop()


if __name__ == "__main__":
    main()
